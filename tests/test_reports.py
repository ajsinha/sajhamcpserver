# Copyright All rights Reserved 2025-2030, Ashutosh Sinha, Email: ajsinha@gmail.com
"""Reports, phase 8.5 core (docs/architecture/Reports.md §1-§4, §8, §12 steps 1-2): definitions and
their load-time checks, runs as the caller (a refused tool is a section, not a crash), report data
and quantities, history and visibility, the formats with marking and provenance, compare, the
generated report_<name> tools and reserved names. Fake tools only; no network."""

import json
import shutil
from pathlib import Path

import pytest

from sajha.observability import caller as caller_mod
from sajha.reports import compare, engine, history, model, render, service, store, tools as rtools
from sajha.tools.base_mcp_tool import BaseMCPTool

ROOT = Path(__file__).resolve().parent.parent


# ── fakes ────────────────────────────────────────────────────────────

class FakeTool(BaseMCPTool):
    def __init__(self, name, fn, schema=None):
        super().__init__({'name': name, 'description': f'fake {name}',
                          'inputSchema': schema or {'type': 'object'}})
        self.fn = fn
        self.calls = []

    def get_input_schema(self):
        return self._input_schema

    def get_output_schema(self):
        return {}

    def execute(self, arguments):
        self.calls.append(dict(arguments))
        return self.fn(arguments)


class FakeRegistry:
    def __init__(self, *tools):
        self.tools = {t.name: t for t in tools}
        self.tool_configs = {}

    def get_tool(self, name):
        return self.tools.get(name)


STATE = {'rate': 4.25}


def _rates(args):
    return {'source': 'BIS WS_CBPOL (CC BY 4.0); cite BIS', 'as_of': '2026-10-10',
            'series': [{'country': 'US', 'title': 'Fed funds', 'observations': [{'date': '2026-09', 'value': STATE['rate']}]},
                       {'country': 'GB', 'title': 'Bank Rate', 'observations': [{'date': '2026-09', 'value': 4.0}]}]}


def _price(args):
    return {'price': 1039.91 + args.get('bump', 0), 'yield': 4.5}


def _boom(args):
    raise RuntimeError('upstream exploded')


def _denied_by_policy(args):
    from sajha.policy.errors import PolicyDenied
    raise PolicyDenied('rule blocks it', tool='policy_tool')


DEF = """
name: rates_pack
title: Rates pack
description: Policy rates and a bond.
owner: alice
parameters:
  countries: {type: array, items: {type: string}, default: [US, GB], description: Areas}
  bump: {type: number, default: 0, description: Price bump}
sections:
  - id: rates
    kind: table
    title: Policy rates
    call: {tool: fake_rates, arguments: {countries: $params.countries}}
    rows: $.series[*]
    key: country
    columns:
      - {key: country, label: Area}
      - {key: latest, label: Latest, path: "$.observations[-1].value", format: percent}
    highlight: {column: latest, above: 4.1}
  - id: bond
    kind: kpis
    call: {tool: fake_price, arguments: {bump: $params.bump, first: "$sections.rates.series[0].country", day: $run.date}}
    values:
      - {label: Price, path: $.price, format: currency}
  - id: secret
    kind: kpis
    call: {tool: fake_secret}
    values: [{label: Secret, path: $.x}]
  - id: missing
    kind: table
    call: {tool: no_such_tool}
    rows: $[*]
  - id: broken
    kind: kpis
    call: {tool: fake_boom}
    values: [{label: X, path: $.x}]
  - id: policy
    kind: kpis
    call: {tool: policy_tool}
    values: [{label: X, path: $.x}]
  - id: curve
    kind: chart
    call: {tool: fake_rates, arguments: {countries: [US]}}
    rows: $.series[*]
    columns: [{key: country, label: Area}, {key: v, label: Rate, path: "$.observations[0].value", format: percent}]
    x: country
    y: v
  - id: commentary
    kind: text
    narrate: {sections: [rates, bond]}
  - id: note
    kind: text
    text: "Areas: ${params.bump}"
formats: [markdown, html]
classification: confidential
"""


@pytest.fixture
def env(tmp_path, monkeypatch):
    import sajha.core.storage as storage
    monkeypatch.setattr(storage, '_storage', storage.LocalStorageBackend(str(tmp_path)))
    (tmp_path / 'config' / 'reports').mkdir(parents=True)
    (tmp_path / 'config' / 'reports' / 'rates_pack.yaml').write_text(DEF)
    monkeypatch.setenv('SAJHA_REPORTS_RELOAD_SECONDS', '0')
    store.reset()
    store.set_tool_exists(lambda n: n != 'no_such_tool')
    reg = FakeRegistry(FakeTool('fake_rates', _rates), FakeTool('fake_price', _price),
                       FakeTool('fake_secret', lambda a: {'x': 42}), FakeTool('fake_boom', _boom),
                       FakeTool('policy_tool', _denied_by_policy))
    engine.set_registry(reg)
    STATE['rate'] = 4.25
    yield tmp_path, reg
    engine.set_registry(None)
    store.set_tool_exists(None)
    store.reset()


def as_user(user, denied=(), admin=False):
    c = caller_mod.Caller(user, '', ('analyst',), 'session', lambda n: n not in denied, admin)
    return caller_mod.set_caller(c)


# ── definitions ──────────────────────────────────────────────────────

def _bad(doc_text):
    with pytest.raises(model.ReportDefinitionError) as e:
        model.parse(doc_text)
    return ' '.join(e.value.problems)


MINIMAL = "name: {name}\nsections: [{{id: a, kind: kpis, call: {{tool: t}}, values: [{{label: X, path: $.x}}]}}]\n"


def test_definition_checks():
    d = model.parse(DEF)
    assert d.tool_name == 'report_rates_pack' and d.formats[0] == 'markdown'
    assert {c['tool'] for c in d.calls()} >= {'fake_rates', 'fake_price', 'no_such_tool'}
    assert 'formats' in d.input_schema()['properties'] and d.input_schema()['examples'] == [{'countries': ['US', 'GB'], 'bump': 0}]
    for reserved in sorted(model.RESERVED):
        assert 'reserved' in _bad(MINIMAL.format(name=reserved))
    assert '"__"' in _bad(MINIMAL.format(name='a__b'))
    assert 'lowercase' in _bad(MINIMAL.format(name='Bad-Name'))
    base = MINIMAL.format(name='ok')
    assert 'names no declared parameter' in _bad(base.replace('call: {tool: t}', 'call: {tool: t, arguments: {x: $params.nope}}'))
    assert 'earlier section' in _bad(base.replace('call: {tool: t}', 'call: {tool: t, arguments: {x: $sections.later.y}}'))
    assert 'path' in _bad(base.replace('$.x', '"$.x[?(@.a)]"'))
    assert 'unknown format' in _bad(base + 'formats: [rtf]\n')
    assert 'recipients' in _bad(base + 'recipients: [{kind: fax}]\n')
    assert 'reserved' in _bad(base + 'parameters: {formats: {type: string}}\n')
    assert 'kind' in _bad(base.replace('kind: kpis', 'kind: pie'))
    assert 'classification' in _bad(base + 'classification: secret\n')
    assert 'max_sections' in str(pytest.raises(model.ReportDefinitionError, model.parse, base, '', 0).value)
    # sub-reports: one level only, never itself
    a = model.parse('name: a\nsections: [{id: s, kind: report, report: {name: b}}]\n')
    b = model.parse('name: b\nsections: [{id: s, kind: report, report: {name: a}}]\n')
    assert set(model.check_sub_reports({'a': a, 'b': b})) == {'a', 'b'}
    assert 'itself' in _bad('name: c\nsections: [{id: s, kind: report, report: {name: c}}]\n')


def test_store_refuses_bad_files_with_a_notice_and_hot_reloads(env, monkeypatch):
    tmp, _ = env
    raised = []
    import sajha.notices as N
    monkeypatch.setattr(N, 'raise_notice', lambda nid, **kw: raised.append((nid, kw.get('severity'))))
    (tmp / 'config' / 'reports' / 'bad.yaml').write_text('name: list\nsections: []\n')
    (tmp / 'config' / 'reports' / 'dup.yaml').write_text(DEF)
    store.scan(force=True)
    assert 'rates_pack' not in store.definitions()           # defined twice: neither loads
    assert any(n == 'reports.definition:bad.yaml' and s == 'warning' for n, s in raised)
    assert any(n == 'reports.duplicate:rates_pack' for n, _ in raised)
    (tmp / 'config' / 'reports' / 'dup.yaml').unlink()
    v = store.version()
    store.scan(force=True)
    assert 'rates_pack' in store.definitions() and store.version() > v - 1
    assert store.missing() == {'rates_pack': ['no_such_tool']}
    assert any(n == 'reports.missing_tools:rates_pack' and s == 'info' for n, s in raised)


def test_shipped_definitions_load_and_their_tools_exist():
    store.reset()
    names = {p.stem for p in (ROOT / 'config' / 'tools').glob('*/*.json')}
    store.set_tool_exists(lambda n: n in names)
    try:
        defs = store.definitions()
        assert {'desk_calendar', 'central_bank_watch'} <= set(defs) and not store.errors()
        assert not store.missing()
        assert defs['desk_calendar'].enabled and not defs['central_bank_watch'].enabled
    finally:
        store.set_tool_exists(None)
        store.reset()


# ── runs ─────────────────────────────────────────────────────────────

def test_run_as_the_caller_sections_fail_alone(env):
    _, reg = env
    t = as_user('bob', denied={'fake_secret'})
    try:
        data = engine.run_report(store.get('rates_pack'), {'bump': 1})
    finally:
        caller_mod.reset(t)
    st = {s['id']: s['status'] for s in data['sections']}
    assert st == {'rates': 'ok', 'bond': 'ok', 'secret': 'not_permitted', 'missing': 'unavailable',
                  'broken': 'error', 'policy': 'not_permitted', 'curve': 'ok', 'commentary': 'pending_narration',
                  'note': 'text'}
    assert data['status'] == 'partial' and len(data['warnings']) == 4
    assert reg.tools['fake_secret'].calls == []               # never called
    sec = {s['id']: s for s in data['sections']}
    assert 'upstream exploded' in sec['broken']['error'] and 'policy' in sec['policy']['error']
    assert sec['secret']['error'] == 'the caller may not run fake_secret'
    # references: parameters, an earlier section, the run date
    call = reg.tools['fake_price'].calls[0]
    assert call['bump'] == 1 and call['first'] == 'US' and call['day'] == data['started_at'][:10]
    assert sec['note']['text'] == 'Areas: 1'
    assert sec['bond']['kpis'][0]['display'] == '1,040.91'
    assert sec['rates']['table']['rows'] == [{'country': 'US', 'latest': 4.25}, {'country': 'GB', 'latest': 4.0}]
    assert sec['rates']['table']['highlight']['rows'] == [0]
    assert sec['rates']['source']['attribution'].startswith('BIS') and sec['rates']['source']['as_of'] == '2026-10-10'
    assert sec['curve']['chart'] == {'type': 'bar', 'x': 'country', 'y': ['v']}
    assert data['runner']['user_id'] == 'bob' and data['classification'] == 'confidential'
    assert data['parameters'] == {'countries': ['US', 'GB'], 'bump': 1}
    # stored, with Markdown and HTML
    stored = history.load(data['run_id'])
    assert stored['run_id'] == data['run_id'] and set(stored['files']) == {'markdown', 'html'}


def test_bad_parameters_disabled_and_roles_are_refused(env):
    d = store.get('rates_pack')
    with pytest.raises(engine.ReportRunError, match='invalid parameters'):
        engine.run_report(d, {'bump': 'x'})
    with pytest.raises(engine.ReportRunError, match='invalid parameters'):
        engine.run_report(d, {'nope': 1})
    d.roles = ['risk']
    t = as_user('bob')
    try:
        with pytest.raises(engine.ReportRunError, match='roles'):
            engine.run_report(d, {})
    finally:
        caller_mod.reset(t)
    d.roles, d.enabled = [], False
    with pytest.raises(engine.ReportRunError, match='disabled'):
        engine.run_report(d, {})


def test_quantities_changes_history_and_visibility(env):
    d = store.get('rates_pack')
    t = as_user('bob')
    try:
        first = engine.run_report(d, {})
        STATE['rate'] = 4.5
        second = engine.run_report(d, {})
        other = engine.run_report(d, {'bump': 2})            # other parameters: not "previous"
    finally:
        caller_mod.reset(t)
    assert second['previous_run_id'] == first['run_id'] and other['previous_run_id'] is None
    q = {x['id']: x for x in second['quantities']}
    assert q['bond.price']['kind'] == 'value' and q['bond.price']['decimals'] == 2
    assert q['bond.price.change']['value'] == 0
    assert q['rates.r0.latest']['value'] == 4.5 and q['rates.r0.latest.rank']['value'] == 1
    assert round(q['rates.r0.latest.share']['value'], 4) == round(4.5 / 8.5 * 100, 4)
    bob = caller_mod.Caller('bob')
    assert [r['run_id'] for r in history.runs('rates_pack', bob)] == [other['run_id'], second['run_id'], first['run_id']]
    assert history.runs('rates_pack', caller_mod.Caller('eve')) == []
    assert len(history.runs('rates_pack', caller_mod.Caller('root', is_admin=True))) == 3
    assert history.runs('rates_pack', bob, status='failed') == []
    assert history.resolve('rates_pack', first['started_at'][:10], bob)['run_id'] == other['run_id']
    assert history.resolve('rates_pack', first['run_id'], caller_mod.Caller('eve')) is None


def test_compare_runs(env):
    d = store.get('rates_pack')
    a = engine.run_report(d, {}, store=False)
    STATE['rate'] = 4.5
    b = engine.run_report(d, {}, store=False)
    b['sections'][0]['table']['rows'].append({'country': 'JP', 'latest': 0.5})
    out = compare.compare_runs(a, b)
    rates = next(s for s in out['sections'] if s['id'] == 'rates')
    assert rates['rows']['key_column'] == 'country'
    assert rates['rows']['added'] == [{'country': 'JP', 'latest': 0.5}] and rates['rows']['removed'] == []
    assert rates['rows']['changed'] == [{'key': 'US', 'cells': {'latest': {'a': 4.25, 'b': 4.5, 'change': 0.25}}}]
    assert rates['rows']['unchanged'] == 1 and out['same_definition']


def test_renderers_marking_watermark_and_provenance(env):
    data = engine.run_report(store.get('rates_pack'), {}, store=False)
    md = render.to_markdown(data)
    assert md.startswith('*CONFIDENTIAL: SAJHA report rates_pack') and md.rstrip().endswith('*')
    assert '## Provenance' in md and 'BIS WS_CBPOL (CC BY 4.0); cite BIS' in md
    assert 'Not permitted: blocked by policy (denied)' in md
    assert 'Unavailable: the tool no_such_tool' in md and 'Commentary not generated yet' in md
    html = render.to_html(data)
    assert 'class="wm"' in html and '<svg' in html and 'CONFIDENTIAL' in html and 'BIS WS_CBPOL' in html
    data['classification'] = 'internal'
    assert 'class="wm"' not in render.to_html(data)
    data['sections'][0]['title'] = '<script>alert(1)</script>'
    assert '<script>alert' not in render.to_html(data)
    assert render.unavailable('markdown') is None and render.unavailable('rtf')


@pytest.mark.parametrize('fmt,module,magic', [('xlsx', 'openpyxl', b'PK'), ('pptx', 'pptx', b'PK'),
                                              ('docx', 'docx', b'PK')])
def test_office_formats(env, fmt, module, magic):
    pytest.importorskip(module, reason=f'{fmt} needs {module}, which is not installed')
    data = engine.run_report(store.get('rates_pack'), {}, store=False)
    out = render.render(data, fmt)
    assert out[:2] == magic and len(out) > 1000
    import io
    if fmt == 'xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(out))
        assert wb.sheetnames[0] == 'Summary' and wb.sheetnames[-1] == 'Provenance' and 'Policy rates' in wb.sheetnames
        assert wb['Summary']['A1'].value.startswith('CONFIDENTIAL')
        assert wb['Summary'].oddHeader.center.text.startswith('CONFIDENTIAL')
    if fmt == 'docx':
        from docx import Document
        doc = Document(io.BytesIO(out))
        assert doc.sections[0].header.paragraphs[0].text.startswith('CONFIDENTIAL')
        assert any(p.text == 'Provenance' for p in doc.paragraphs)


def test_missing_office_package_is_unavailable_with_a_reason(monkeypatch):
    monkeypatch.setattr(render, '_has', lambda m: False)
    assert 'python-pptx' in render.unavailable('pptx') and 'openpyxl' in render.unavailable('xlsx')
    data = {'run_id': 'x-20260101T000000000-abcdef', 'files': {}, 'files_unavailable': {}}
    engine.render_and_store(data, ['pptx'])
    assert 'pptx' not in data['files'] and 'python-pptx' in data['files_unavailable']['pptx']


def test_pdf(env):
    if render.unavailable('pdf'):
        pytest.skip(render.unavailable('pdf'))
    data = engine.run_report(store.get('rates_pack'), {}, store=False)
    pdf = render.render(data, 'pdf')
    assert pdf[:5] == b'%PDF-' and len(pdf) > 2000


# ── tools ────────────────────────────────────────────────────────────

def test_report_tools_end_to_end(env):
    t = as_user('bob')
    try:
        listed = rtools.ReportListTool({'name': 'report_list'}).execute({'search': 'rates'})
        assert [r['name'] for r in listed['reports']] == ['rates_pack'] and 'formats' not in listed['reports'][0]['parameters']
        desc = rtools.ReportDescribeTool({'name': 'report_describe'}).execute({'report': 'rates_pack'})
        assert desc['missing_tools'] == ['no_such_tool'] and desc['tool'] == 'report_rates_pack'
        gen = rtools.ReportTool(service.tool_config(store.get('rates_pack')))
        out = gen.execute({'bump': 0, 'formats': []})
        assert out['status'] == 'partial' and out['markdown'].startswith('*CONFIDENTIAL') and out['files'] == {'markdown': out['files']['markdown']}
        run2 = rtools.ReportRunTool({'name': 'report_run'}).execute({'report': 'rates_pack', 'parameters': {}})
        hist = rtools.ReportHistoryTool({'name': 'report_history'}).execute({'report': 'rates_pack'})
        assert [r['run_id'] for r in hist['runs']] == [run2['run_id'], out['run_id']]
        r = rtools.ReportRenderTool({'name': 'report_render'}).execute({'run_id': out['run_id'], 'format': 'html'})
        assert r['path'].endswith('report.html') and r['bytes'] > 500
        assert 'html' in history.load(out['run_id'])['files']
        cmp_ = rtools.ReportCompareTool({'name': 'report_compare'}).execute(
            {'report': 'rates_pack', 'run_a': out['run_id'], 'run_b': run2['run_id']})
        assert cmp_['a']['run_id'] == out['run_id']
        with pytest.raises(ValueError, match='no report'):
            rtools.ReportRunTool({'name': 'report_run'}).execute({'report': 'nope'})
    finally:
        caller_mod.reset(t)
    t = as_user('eve')
    try:
        with pytest.raises(ValueError, match='no run'):
            rtools.ReportRenderTool({'name': 'report_render'}).execute({'run_id': out['run_id'], 'format': 'html'})
    finally:
        caller_mod.reset(t)
    t = as_user('mallory', denied={'report_rates_pack'})
    try:
        assert rtools.ReportListTool({'name': 'report_list'}).execute({})['reports'] == []
        with pytest.raises(PermissionError):
            rtools.ReportRunTool({'name': 'report_run'}).execute({'report': 'rates_pack'})
    finally:
        caller_mod.reset(t)


def test_a_report_calling_itself_is_refused_as_a_section(env):
    _, reg = env
    tool = rtools.ReportTool(service.tool_config(store.get('rates_pack')))
    reg.tools['report_rates_pack'] = tool
    (env[0] / 'config' / 'reports' / 'loop.yaml').write_text(
        'name: loopy\nsections: [{id: a, kind: kpis, call: {tool: report_loopy}, values: [{label: X, path: $.x}]}]\n')
    store.scan(force=True)
    loop_tool = rtools.ReportTool(service.tool_config(store.get('loopy')))
    reg.tools['report_loopy'] = loop_tool
    out = loop_tool.execute({'formats': []})
    sec = out['sections'][0]
    assert out['status'] == 'failed' and sec['status'] in ('not_permitted', 'error')


def test_generated_configs_sync_and_lint(env):
    from sajha.quality.lint import lint_tool
    cfg = service.tool_config(store.get('rates_pack'))
    assert cfg['name'] == 'report_rates_pack' and cfg['report']['name'] == 'rates_pack'
    assert cfg['annotations']['readOnlyHint'] is True and len(cfg['description']) <= service.MAX_DESCRIPTION
    findings = [f for f in lint_tool(cfg['name'], cfg) if f.level in ('error', 'warning')]
    assert not findings, findings
    tmp, _ = env
    out = service.sync(None)
    assert out['added'] == ['report_rates_pack']
    written = tmp / 'config' / 'tools' / 'reports' / 'report_rates_pack.json'
    assert json.loads(written.read_text())['report']['name'] == 'rates_pack'
    assert service.sync(None)['unchanged'] == ['report_rates_pack']
    (tmp / 'config' / 'reports' / 'rates_pack.yaml').write_text(DEF.replace('owner: alice', 'owner: alice\nenabled: false'))
    store.scan(force=True)
    assert service.sync(None)['removed'] == ['report_rates_pack'] and not written.exists()


def test_name_taken_by_another_tool_is_a_conflict(env):
    tmp, _ = env
    other = tmp / 'config' / 'tools' / 'misc'
    other.mkdir(parents=True)
    (other / 'report_rates_pack.json').write_text(json.dumps({'name': 'report_rates_pack', 'implementation': 'x.Y'}))
    from sajha.tools import tool_files
    tool_files.get_index().refresh()
    out = service.sync(None)
    assert out['conflicts'] == ['report_rates_pack'] and out['added'] == []
    assert json.loads((other / 'report_rates_pack.json').read_text())['implementation'] == 'x.Y'


def test_shipped_report_tool_configs_lint_clean():
    from sajha.quality.lint import lint_tool
    from tests.tool_config_files import read_tool
    for name in ('report_list', 'report_describe', 'report_run', 'report_history', 'report_render', 'report_compare'):
        cfg = read_tool(name)
        assert cfg['implementation'].startswith('sajha.reports.tools.')
        bad = [f for f in lint_tool(name, cfg) if f.level in ('error', 'warning')]
        assert not bad, bad
    assert (ROOT / 'config' / 'tools' / 'reports' / 'domain.txt').read_text().strip()


def test_flagship_offline_report_runs_on_the_shipped_tools(tmp_path, monkeypatch):
    pytest.importorskip('holidays')
    import importlib
    import sajha.core.storage as storage
    from tests.tool_config_files import read_tool
    monkeypatch.setattr(storage, '_storage', storage.LocalStorageBackend(str(tmp_path)))
    shutil.copytree(ROOT / 'config' / 'reports', tmp_path / 'config' / 'reports')
    store.reset()
    store.set_tool_exists(lambda n: True)

    class Reg:
        def get_tool(self, n):
            cfg = read_tool(n)
            m, c = cfg['implementation'].rsplit('.', 1)
            return getattr(importlib.import_module(m), c)(cfg)
    engine.set_registry(Reg())
    try:
        data = engine.run_report(store.get('desk_calendar'), {'market': 'XNYS'}, formats=['markdown', 'html'])
    finally:
        engine.set_registry(None)
        store.set_tool_exists(None)
        store.reset()
    assert data['status'] == 'succeeded', data['warnings']
    sec = {s['id']: s for s in data['sections']}
    assert sec['holidays']['table']['row_count'] >= 9 and sec['bond']['kpis'][0]['display'] == '1,039.91'
    assert 'holidays Python package' in sec['holidays']['source']['attribution']


def test_runs_and_renders_are_audited(env, monkeypatch):
    from sajha import audit
    seen = []
    monkeypatch.setattr(audit, 'record', lambda event, **kw: seen.append((event, kw)))
    t = as_user('bob')
    try:
        data = engine.run_report(store.get('rates_pack'), {}, formats=[])
        rtools.ReportRenderTool({'name': 'report_render'}).execute({'run_id': data['run_id'], 'format': 'html'})
    finally:
        caller_mod.reset(t)
    runs = [kw for e, kw in seen if e == 'report.run']
    assert runs and runs[0]['resource'] == {'type': 'report', 'id': 'rates_pack', 'run_id': data['run_id']}
    assert runs[0]['actor']['user'] == 'bob' and runs[0]['outcome'] == 'partial'
    assert [kw['details']['format'] for e, kw in seen if e == 'report.render'] == ['html']


def test_absolute_dirs_outside_the_storage_base(env, tmp_path_factory, monkeypatch):
    outside = tmp_path_factory.mktemp('outside')
    (outside / 'defs').mkdir()
    (outside / 'defs' / 'rates_pack.yaml').write_text(DEF)
    monkeypatch.setenv('SAJHA_REPORTS_DIR', str(outside / 'defs'))
    monkeypatch.setenv('SAJHA_REPORTS_OUTPUT_DIR', str(outside / 'runs'))
    store.reset()
    assert 'rates_pack' in store.definitions()
    t = as_user('bob')
    try:
        data = engine.run_report(store.get('rates_pack'), {}, formats=['html'])
    finally:
        caller_mod.reset(t)
    assert (outside / 'runs' / 'rates_pack' / data['run_id'] / 'report.html').exists()
    assert [r['run_id'] for r in history.runs('rates_pack', caller_mod.Caller('bob'))] == [data['run_id']]
