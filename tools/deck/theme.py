"""
The deck design system: SAJHA Crimson, the light theme of the web UI, Georgia for
headings and Calibri for text, and the layout primitives every slide is drawn with.

The colours are not restated here. They are read from the one place SAJHA defines
colour, the first ``:root`` block of ``sajha/web/static/css/tokens.css`` (the Crimson
theme), so the deck and the product cannot drift apart; a token that disappears from
that file fails the build.

Every primitive that places text measures it with ``metrics`` — the same estimator the
geometry audit uses — and the fitting helpers shrink the type or refuse, so a slide
that would overflow fails the build rather than the reader.

Copyright All rights Reserved 2025-2030, Ashutosh Sinha, Email: ajsinha@gmail.com
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Callable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt
from pptx.util import Inches as In

from metrics import FOOTER_Y, SAFETY, SANS, SERIF, SH, SW, est_lines, text_extent, text_h

TOKENS = Path(__file__).resolve().parents[2] / "sajha" / "web" / "static" / "css" / "tokens.css"


def _tokens(path: Path = TOKENS) -> dict[str, RGBColor]:
    """``--sajha-<name>: #RRGGBB`` from the first ``:root`` block (the Crimson theme)."""
    css = path.read_text(encoding="utf-8")
    block = css[css.index(":root {") : css.index("}", css.index(":root {"))]
    found = re.findall(r"--sajha-([a-z-]+):\s*#([0-9A-Fa-f]{6})\b", block)
    return {name: RGBColor.from_string(hex_.upper()) for name, hex_ in found}


_T = _tokens()

CRIMSON = _T["crimson"]  # the accent
CRIMSON_D = _T["crimson-deep"]  # headings, dividers' edge
CRIMSON_S = _T["crimson-strong"]
PINK = _T["glow"]  # light text on crimson
PINK_L = _T["crimson-tint"]
INK = _T["ink"]
SLATE = _T["slate"]
MUTED = _T["slate"]
RULE = _T["border"]
PARCH = _T["canvas"]
WHITE = _T["surface"]
GOLD = _T["highlight"]
NAVY = _T["indigo"]
NAV_FROM = _T["nav-from"]
OK = _T["ok"]
WARN = _T["warn"]
BAD = _T["bad"]


def _tint(c: RGBColor, k: float) -> RGBColor:
    """c moved a fraction k of the way to white."""
    return RGBColor(*(round(v + (255 - v) * k) for v in c))


WATERMARK = _tint(CRIMSON, 0.15)  # a chapter numeral one tone lighter than a divider

ML = 0.85
CW = SW - 2 * ML
BODY_BOTTOM = FOOTER_Y - 0.12

_state: dict[str, Any] = {"prs": None, "chapter": "", "n": 0}


class DoesNotFit(Exception):
    """A block that cannot be made to fit its box at the smallest permitted size."""


def new_deck(chapter: str) -> Any:
    """Start a fresh 16:9 presentation and make it the current one."""
    prs = Presentation()
    prs.slide_width = In(SW)
    prs.slide_height = In(SH)
    _state.update(prs=prs, chapter=chapter, n=0)
    return prs


def chapter(name: str) -> None:
    _state["chapter"] = name


def blank() -> Any:
    prs = _state["prs"]
    return prs.slides.add_slide(prs.slide_layouts[6])


def remove(shape: Any) -> None:
    el = shape._element
    el.getparent().remove(el)


def rect(
    sl: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: Any = None,
    line: Any = None,
    lw: float = 1.0,
    shape: Any = MSO_SHAPE.RECTANGLE,
) -> Any:
    s = sl.shapes.add_shape(shape, In(x), In(y), In(w), In(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s


def txt(
    sl: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    align: Any = PP_ALIGN.LEFT,
    anchor: Any = MSO_ANCHOR.TOP,
) -> Any:
    tb = sl.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def para(
    tf: Any,
    text: str,
    size: float = 14,
    color: Any = INK,
    bold: bool = False,
    font: str = SANS,
    italic: bool = False,
    space_after: float = 6,
    first: bool = False,
    line: float = 1.2,
    space_before: float = 0,
) -> Any:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return p


def runs(
    tf: Any,
    parts: list[tuple],
    size: float = 14,
    space_after: float = 6,
    first: bool = False,
    line: float = 1.2,
    level: int = 0,
    space_before: float = 0,
) -> Any:
    """``parts`` is a list of ``(text, color, bold)`` tuples, one run each."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    p.level = level
    for text, color, bold in parts:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = SANS
    return p


def fitted(
    sl: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    write: Callable[[Any, float], None],
    start: float,
    floor: float = 9.0,
) -> float:
    """Write a text block at the largest size in [floor, start] whose estimated
    extent fits ``h``; return the height it uses. Raises ``DoesNotFit``."""
    size = start
    while size >= floor:
        tf = txt(sl, x, y, w, h)
        write(tf, size)
        need = text_extent(sl.shapes[-1])
        if need <= h:
            return need
        remove(sl.shapes[-1])
        size -= 0.5
    raise DoesNotFit(f"text block does not fit {w:.2f}x{h:.2f} at {floor}pt")


FOOTER_TEXT = "SAJHA \u2022 Ashutosh Sinha"
ICONS = TOKENS.parents[1] / "icons" / "sajha-icons.svg"


def _mark_geometry(path: Path = ICONS) -> tuple[float, list[tuple[float, float, float]], list[tuple[float, ...]]]:
    """SAJHA's own mark, the ``icon-sajha`` symbol the web console's brand uses: its
    view box, circles and lines, read from the icon sprite so the deck draws the same mark."""
    svg = path.read_text(encoding="utf-8")
    m = re.search(r'<symbol id="icon-sajha" viewBox="0 0 (\d+) \d+"[^>]*>(.*?)</symbol>', svg, re.S)
    if not m:
        raise ValueError(f"no icon-sajha symbol in {path}")
    body = m.group(2)
    circles = [tuple(float(v) for v in c) for c in re.findall(r'<circle cx="([\d.]+)" cy="([\d.]+)" r="([\d.]+)"', body)]
    lines = [tuple(float(v) for v in ln) for ln in re.findall(
        r'<line x1="([\d.]+)" y1="([\d.]+)" x2="([\d.]+)" y2="([\d.]+)"', body)]
    if not circles or not lines:
        raise ValueError("the icon-sajha symbol has no circles or lines")
    return float(m.group(1)), circles, lines  # type: ignore[return-value]


MARK = _mark_geometry()


def mark(sl: Any, x: float, y: float, size: float, color: Any, lw: float = 1.5) -> None:
    """Draw SAJHA's mark (three joined circles) in a ``size`` square at (x, y)."""
    box, circles, lines = MARK
    k = size / box
    for x1, y1, x2, y2 in lines:
        c = sl.shapes.add_connector(1, In(x + x1 * k), In(y + y1 * k), In(x + x2 * k), In(y + y2 * k))
        c.line.color.rgb = color
        c.line.width = Pt(lw)
    for cx, cy, r in circles:
        o = sl.shapes.add_shape(MSO_SHAPE.OVAL, In(x + (cx - r) * k), In(y + (cy - r) * k), In(2 * r * k), In(2 * r * k))
        o.fill.background()
        o.line.color.rgb = color
        o.line.width = Pt(lw)
        o.shadow.inherit = False


def footer(sl: Any, color: Any = None) -> None:
    color = color or MUTED
    rect(sl, ML, FOOTER_Y, CW, 0.008, fill=RULE if color is MUTED else color)
    tf = txt(sl, ML, SH - 0.46, CW * 0.7, 0.24)
    para(tf, FOOTER_TEXT, size=8.5, color=color, first=True, space_after=0)
    tf = txt(sl, ML + CW * 0.7, SH - 0.46, CW * 0.3, 0.24, align=PP_ALIGN.RIGHT)
    para(tf, str(_state["n"]), size=8.5, color=color, first=True, space_after=0, bold=True)


def brand(sl: Any, color: Any = None, light: bool = False) -> None:
    """The mark and the wordmark in the top-right corner of a content slide."""
    color = color or CRIMSON
    x = SW - ML - 1.05
    mark(sl, x, 0.16, 0.30, color, 1.1)
    tf = txt(sl, x + 0.38, 0.17, 0.75, 0.28)
    para(tf, "SAJHA", size=11, color=color, bold=True, font=SERIF, first=True, space_after=0)


def content(title: str, kicker: str | None = None) -> tuple[Any, float]:
    """A content slide's chrome. Returns ``(slide, body_top)``."""
    _state["n"] += 1
    sl = blank()
    rect(sl, 0, 0, SW, SH, fill=WHITE)
    rect(sl, 0, 0.62, 0.30, 0.055, fill=CRIMSON)
    y = 0.52
    if kicker:
        tf = txt(sl, ML, y, CW, 0.24)
        para(tf, kicker.upper(), size=10.5, color=CRIMSON, bold=True, first=True, space_after=0)
        y += 0.30
    size = 26.0
    while size > 20 and est_lines(title, CW * SAFETY, size, False, SERIF) > 1:
        size -= 1
    th = text_h(title, CW * SAFETY, size, False, SERIF, 1.15)
    tf = txt(sl, ML, y, CW, th + 0.04)
    para(tf, title, size=size, color=INK, font=SERIF, first=True, space_after=0, line=1.15)
    body_top = y + th + 0.24
    rect(sl, ML, body_top - 0.14, CW, 0.012, fill=RULE)
    brand(sl)
    footer(sl)
    return sl, body_top


SECTION_SHAPE = "Section numeral"


def divider(num: str, title: str, sub: str, points: list[str]) -> Any:
    """A numbered section divider: the chapter number as a large watermark, the section's title, one
    sentence on what it answers, and its contents in a line."""
    _state["chapter"] = f"{num} · {title}"
    _state["n"] += 1
    sl = blank()
    rect(sl, 0, 0, SW, SH, fill=CRIMSON)
    rect(sl, 0, 0, 0.18, SH, fill=CRIMSON_D)
    mark(sl, ML, 0.42, 0.42, PINK, 1.4)
    tf = txt(sl, ML + 0.55, 0.47, 3.0, 0.32)
    para(tf, "SAJHA", size=14, color=PINK, bold=True, font=SERIF, first=True, space_after=0)
    # The chapter number as a watermark: two digits, very large, one tone lighter than the
    # background, behind everything else on the slide (drawn before the title).
    wm = txt(sl, SW - ML - 4.6, 0.05, 4.6 + ML - 0.25, 2.75, align=PP_ALIGN.RIGHT)
    wm.margin_left = wm.margin_right = wm.margin_top = wm.margin_bottom = 0
    wm.vertical_anchor = MSO_ANCHOR.TOP
    para(wm, f"{int(num):02d}" if num.isdigit() else num, size=170, color=WATERMARK, font=SANS,
         bold=True, first=True, space_after=0, line=1.0)
    sl.shapes[-1].name = SECTION_SHAPE
    tw = CW * 0.86
    size = 40.0
    while size > 26 and est_lines(title, tw * SAFETY, size, False, SERIF) > 1:
        size -= 2
    th = text_h(title, tw * SAFETY, size, False, SERIF, 1.1)
    tf = txt(sl, (SW - tw) / 2, 2.72, tw, th + 0.05, align=PP_ALIGN.CENTER)
    para(tf, title, size=size, color=WHITE, font=SERIF, first=True, space_after=0, line=1.1)
    y = 2.72 + th + 0.22
    rect(sl, (SW - 1.6) / 2, y, 1.6, 0.035, fill=PINK)
    y += 0.25
    sw = CW * 0.74

    def write_sub(tf: Any, s: float) -> None:
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        para(tf, sub, size=s, color=PINK_L, italic=True, first=True, space_after=0, line=1.25)

    used = fitted(sl, (SW - sw) / 2, y, sw, 1.45, write_sub, 16, 11)
    y += used + 0.32

    def write_points(tf: Any, s: float) -> None:
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        para(tf, "   ·   ".join(points), size=s, color=PINK, bold=True, first=True, space_after=0, line=1.3)

    fitted(sl, ML + 0.4, y, CW - 0.8, FOOTER_Y - 0.15 - y, write_points, 12.5, 9)
    footer(sl, PINK)
    return sl


def arrow(sl: Any, x1: float, y1: float, x2: float, y2: float, color: Any = None, width: float = 1.5) -> Any:
    """A straight connector with an arrowhead at (x2, y2)."""
    from pptx.oxml.ns import qn

    c = connect(sl, x1, y1, x2, y2, color or SLATE, width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def boxed(
    sl: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    fill: Any = None,
    color: Any = None,
    line: Any = None,
    start: float = 13,
    floor: float = 8.5,
    bold: bool = True,
    shape: Any = MSO_SHAPE.ROUNDED_RECTANGLE,
    align: Any = PP_ALIGN.CENTER,
    sub: str = "",
    font: str = SANS,
) -> Any:
    """A filled shape with its own text, centred, at the largest size that fits; lines in
    ``text`` separated by newlines become paragraphs. ``sub`` is a smaller second line."""
    s = sl.shapes.add_shape(shape, In(x), In(y), In(w), In(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.0)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, 0.08 / max(0.1, min(w, h)))
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = In(0.06)
    tf.margin_top = tf.margin_bottom = In(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    color = color or INK
    size = start
    while True:
        for p in list(tf.paragraphs)[1:]:
            p._p.getparent().remove(p._p)
        tf.paragraphs[0].text = ""
        for r in list(tf.paragraphs[0].runs):
            r._r.getparent().remove(r._r)
        heads = text.split("\n")
        for i, t in enumerate(heads):
            p = para(tf, t, size=size, color=color, bold=bold, first=i == 0, space_after=0, line=1.05, font=font)
            p.alignment = align
        if sub:
            p = para(tf, sub, size=max(7.5, size - 2.5), color=color, first=False, space_after=0, line=1.05)
            p.alignment = align
        if text_extent(s) <= h - 0.06 or size <= floor:
            break
        size -= 0.5
    if text_extent(s) > h - 0.02:
        raise DoesNotFit(f"box text {text!r} does not fit {w:.2f}x{h:.2f} at {floor}pt")
    return s


def numdot(sl: Any, x: float, y: float, d: float, num: str, fill: Any = None, color: Any = None) -> Any:
    """A numbered circle."""
    return boxed(sl, x, y, d, d, num, fill=fill or CRIMSON, color=color or WHITE, start=min(16, d * 30),
                 floor=7, shape=MSO_SHAPE.OVAL)


def _row_heights(
    data: list[list[str]], widths: list[float], fs: float, hfs: float, bold_col0: bool
) -> list[float]:
    heights = []
    for r, row in enumerate(data):
        size = hfs if r == 0 else fs
        need = 0.0
        for c, cell in enumerate(row):
            bold = r == 0 or (bold_col0 and c == 0)
            need = max(need, text_h(cell, (widths[c] - 0.18) * SAFETY, size, bold, SANS, 1.0))
        heights.append(max(0.30, need + 0.12))
    return heights


def _fill_cell(
    cell: Any, text: str, r: int, c: int, fs: float, hfs: float, bold_col0: bool
) -> None:
    cell.margin_left = In(0.09)
    cell.margin_right = In(0.07)
    cell.margin_top = cell.margin_bottom = In(0.035)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.fill.solid()
    cell.fill.fore_color.rgb = CRIMSON if r == 0 else (PARCH if r % 2 == 0 else WHITE)
    tf = cell.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = str(text)
    run.font.name = SANS
    run.font.size = Pt(hfs if r == 0 else fs)
    run.font.bold = r == 0 or (bold_col0 and c == 0)
    run.font.color.rgb = WHITE if r == 0 else (CRIMSON_D if bold_col0 and c == 0 else INK)


def table(
    sl: Any,
    data: list[list[str]],
    x: float,
    y: float,
    w: float,
    col_w: list[float] | None = None,
    fs: float = 11.5,
    hfs: float = 11,
    bold_col0: bool = True,
) -> float:
    """A table whose row heights are computed from its wrapped text.
    Returns the rendered height, so callers place what follows from it."""
    cols = len(data[0])
    col_w = col_w or [1.0] * cols
    widths = [w * c / sum(col_w) for c in col_w]
    heights = _row_heights(data, widths, fs, hfs, bold_col0)
    total = sum(heights)
    gf = sl.shapes.add_table(len(data), cols, In(x), In(y), In(w), In(total))
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    for i, cw in enumerate(widths):
        tbl.columns[i].width = Emu(int(In(cw)))
    for r, row in enumerate(data):
        tbl.rows[r].height = Emu(int(In(heights[r])))
        for c, cell in enumerate(row):
            _fill_cell(tbl.cell(r, c), cell, r, c, fs, hfs, bold_col0)
    return total


def fitted_table(
    sl: Any,
    data: list[list[str]],
    x: float,
    y: float,
    w: float,
    h: float,
    col_w: list[float] | None = None,
    start: float = 14.0,
    bold_col0: bool = True,
) -> float:
    """``table`` at the largest size in [9, start] that fits ``h``."""
    fs = start
    while fs >= 9.0:
        used = table(sl, data, x, y, w, col_w, fs=fs, hfs=min(fs, 15), bold_col0=bold_col0)
        if used <= h:
            return used
        remove(sl.shapes[-1])
        fs -= 0.5
    raise DoesNotFit(f"table of {len(data)} rows does not fit {h:.2f}in")


def card(sl: Any, x: float, y: float, w: float, h: float, num: str, title: str, body: str) -> None:
    """A bordered card whose contents are fitted inside the border."""
    rect(sl, x, y, w, h, fill=WHITE, line=RULE, lw=0.9)
    rect(sl, x, y, w, 0.055, fill=CRIMSON)
    inner = w - 0.40
    top = y + 0.18
    if num:
        tf = txt(sl, x + 0.20, top, inner, 0.24)
        para(tf, num, size=10, color=CRIMSON, bold=True, first=True, space_after=0)
        top += 0.28
    used = fitted(
        sl,
        x + 0.20,
        top,
        inner,
        0.72,
        lambda tf, s: para(
            tf,
            title,
            size=s,
            color=INK,
            bold=True,
            font=SERIF,
            first=True,
            space_after=0,
            line=1.12,
        ),
        19,
        10.5,
    )
    by = top + used + 0.10
    fitted(
        sl,
        x + 0.20,
        by,
        inner,
        (y + h - 0.14) - by,
        lambda tf, s: para(tf, body, size=s, color=SLATE, first=True, space_after=0, line=1.2),
        17,
        8.5,
    )


def statbar(sl: Any, y: float, stats: list[tuple[str, str]], h: float = 1.25) -> None:
    gap = 0.22
    bw = (CW - gap * (len(stats) - 1)) / len(stats)
    for i, (big, label) in enumerate(stats):
        x = ML + i * (bw + gap)
        rect(sl, x, y, bw, h, fill=PARCH)
        rect(sl, x, y, 0.045, h, fill=CRIMSON)
        fitted(
            sl,
            x + 0.22,
            y + 0.14,
            bw - 0.34,
            0.5,
            lambda tf, s, b=big: para(
                tf,
                b,
                size=s,
                color=CRIMSON,
                bold=True,
                font=SERIF,
                first=True,
                space_after=0,
                line=1.0,
            ),
            24,
            14,
        )
        fitted(
            sl,
            x + 0.22,
            y + 0.66,
            bw - 0.34,
            h - 0.74,
            lambda tf, s, lab=label: para(
                tf, lab, size=s, color=SLATE, first=True, space_after=0, line=1.12
            ),
            13,
            8,
        )


def connect(
    sl: Any, x1: float, y1: float, x2: float, y2: float, color: Any = SLATE, width: float = 1.5
) -> Any:
    """A straight connector. Elbow connectors auto-route into detours that
    collide with the nodes they join; do not use them."""
    c = sl.shapes.add_connector(1, In(x1), In(y1), In(x2), In(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def notes(sl: Any, text: str) -> None:
    """The speaker notes: where every number and claim on the slide comes from."""
    sl.notes_slide.notes_text_frame.text = text


def panel(sl: Any, x: float, y: float, w: float, h: float, lines: list[str], start: float = 12.5) -> float:
    """A filled panel of monospaced lines (a captured run, a config excerpt), fitted inside
    its border. Returns the height used by the text."""
    rect(sl, x, y, w, h, fill=PARCH, line=RULE, lw=0.75)
    rect(sl, x, y, 0.045, h, fill=CRIMSON)

    def write(tf: Any, s: float) -> None:
        for i, ln in enumerate(lines):
            p = para(tf, ln, size=s, color=INK, font="Consolas", first=i == 0, space_after=2, line=1.1)
            if ln.startswith("#"):
                p.runs[0].font.color.rgb = SLATE
            elif ln.startswith("→") or ln.startswith("✗") or ln.startswith("✓"):
                p.runs[0].font.color.rgb = CRIMSON_D
                p.runs[0].font.bold = True

    # A monospaced panel is a captured run: a wrapped line misaligns its columns, so the type
    # is never larger than the size at which the longest line fits on one line.
    longest = max((len(ln) for ln in lines), default=1)
    widest = (w - 0.42) * SAFETY * 72.0 / (0.620 * max(longest, 1))
    return fitted(sl, x + 0.25, y + 0.16, w - 0.42, h - 0.30, write, max(8.0, min(start, widest)), 8.0)
