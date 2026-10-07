"""
SAJHA MCP Server — RAG: text from PDF and Word documents.
Copyright All rights Reserved 2025-2030, Ashutosh Sinha

PDF (``.pdf``) and Word (``.docx``) files are indexed beside the text types in chunking.py.
Their readers are optional packages, imported only when such a file is met:

* ``.pdf``  needs ``pypdf``        (``pip install pypdf``)
* ``.docx`` needs ``python-docx``  (``pip install python-docx``)

Without the package the file is not indexed and the index build reports
"install pypdf to index PDF files" (or the Word equivalent); every other document is
indexed as usual. Word headings (styles "Title" and "Heading N") come out as Markdown
headings, so a Word document is split by section like a guide; a PDF is split by
paragraphs, with each page's text in order. Scanned PDFs (images, no text layer) give no
text and are reported as such: there is no OCR.
"""

from __future__ import annotations

import io
import re
from typing import Tuple

PDF_TYPES = (".pdf",)
WORD_TYPES = (".docx",)
DOCUMENT_TYPES = PDF_TYPES + WORD_TYPES
_PACKAGE = {".pdf": ("pypdf", "pypdf", "PDF"), ".docx": ("docx", "python-docx", "Word (.docx)")}


class ExtractorMissing(RuntimeError):
    """The optional package that reads this file type is not installed."""


def is_document(name: str) -> bool:
    """True for a binary document type read through an optional package (PDF, Word)."""
    return (name or "").lower().endswith(DOCUMENT_TYPES)


def _suffix(name: str) -> str:
    low = (name or "").lower()
    return next((s for s in DOCUMENT_TYPES if low.endswith(s)), "")


def require(name: str) -> None:
    """Raise :class:`ExtractorMissing` (with the pip command) when ``name``'s reader is not installed."""
    suffix = _suffix(name)
    if not suffix:
        return
    module, package, label = _PACKAGE[suffix]
    try:
        __import__(module)
    except ImportError:
        raise ExtractorMissing(f"install {package} to index {label} files (pip install {package})") from None


def available() -> dict:
    """{suffix: installed?} for the optional readers (shown in the index stats)."""
    out = {}
    for suffix in DOCUMENT_TYPES:
        try:
            require("x" + suffix)
            out[suffix] = True
        except ExtractorMissing:
            out[suffix] = False
    return out


def extract_text(name: str, data: bytes) -> Tuple[str, str]:
    """``(text, title)`` of a PDF or Word file. ``title`` is "" when the file names none.

    Raises :class:`ExtractorMissing` without the reader, ``ValueError`` for a file it cannot
    read or one with no text."""
    require(name)
    suffix = _suffix(name)
    try:
        text, title = (_pdf if suffix == ".pdf" else _docx)(data)
    except ExtractorMissing:
        raise
    except Exception as e:
        raise ValueError(f"cannot read {name.rsplit('/', 1)[-1]}: {e}") from None
    text = re.sub(r"[ \t]+\n", "\n", text or "").strip()
    if not text:
        raise ValueError(f"{name.rsplit('/', 1)[-1]} has no text to index"
                         + (" (a scanned PDF needs OCR, which SAJHA does not do)" if suffix == ".pdf" else ""))
    return text, (title or "").strip()


def _pdf(data: bytes) -> Tuple[str, str]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        t = (page.extract_text() or "").strip()
        if t:
            pages.append(t)
    title = ""
    try:
        title = str((reader.metadata or {}).get("/Title") or "")
    except Exception:
        pass
    return "\n\n".join(pages), title


def _docx(data: bytes) -> Tuple[str, str]:
    import docx
    doc = docx.Document(io.BytesIO(data))
    lines, title = [], ""
    for p in doc.paragraphs:
        text = (p.text or "").strip()
        if not text:
            continue
        style = (getattr(p.style, "name", "") or "").lower()
        m = re.match(r"heading (\d)", style)
        if style == "title":
            title = title or text
            lines.append(f"# {text}")
        elif m:
            lines.append("#" * min(6, int(m.group(1)) + 1) + " " + text)
        else:
            lines.append(text)
    for table in doc.tables:                 # tables after the body text, one row per line
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    try:
        title = title or (doc.core_properties.title or "")
    except Exception:
        pass
    return "\n\n".join(lines), title
